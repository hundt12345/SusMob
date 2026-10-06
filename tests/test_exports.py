"""Export-Tests: die .xlsx muss echte Formeln und ein Annahmen-Blatt enthalten."""
import io

from openpyxl import load_workbook

from server import calc, exports


def _wb(daten: bytes):
    return load_workbook(io.BytesIO(daten))


def test_excel_co2_mit_formeln_und_annahmen():
    ergebnis = calc.co2_bilanz([
        {"bezeichnung": "Pkw Diesel", "anzahl": 25, "km_jahr": 28000, "verbrauch_je_100km": 6.8,
         "kraftstoff": "diesel"},
        {"bezeichnung": "Busse", "anzahl": 3, "km_jahr": 85000, "verbrauch_je_100km": 45.0,
         "kraftstoff": "diesel"},
    ])
    daten = exports.build_workbook(tile_name="CO₂-Bilanz", conversation_title="Testlauf",
                                   calc=ergebnis, structured={"annahmen": ["Strommix 372 g/kWh"]},
                                   answer="## Ergebnis\n\n| Gruppe | t CO₂ |\n| --- | --- |\n| Busse | 307,5 |",
                                   model="nvidia/test:free")
    wb = _wb(daten)
    assert {"Ergebnis", "Annahmen", "Bilanz", "Hinweise", "Antworttext"} <= set(wb.sheetnames)
    bilanz = wb["Bilanz"]
    # Energie- und Faktor-Spalte sind Formeln, keine festen Werte
    assert str(bilanz["E4"].value).startswith("=")
    assert "VLOOKUP" in str(bilanz["F4"].value)
    assert str(bilanz["G4"].value).startswith("=")
    # Annahmen-Blatt hat editierbare Faktoren
    annahmen = wb["Annahmen"]
    assert any(annahmen.cell(row=r, column=1).value == "diesel" for r in range(1, 30))
    assert "Prüfpflicht" in str(wb["Ergebnis"]["A6"].value)


def test_excel_massnahmen_nutzt_gewichte_aus_annahmen():
    ergebnis = calc.massnahmen_priorisierung([
        {"name": "Fahrradstraße", "wirkung": 4, "sicherheit": 5, "kosten": 4, "aufwand": 4, "akzeptanz": 3},
        {"name": "Stadtbahn", "wirkung": 5, "sicherheit": 3, "kosten": 1, "aufwand": 1, "akzeptanz": 2},
    ])
    daten = exports.build_workbook(tile_name="Maßnahmenplanung", calc=ergebnis, structured={}, answer="")
    wb = _wb(daten)
    assert "Maßnahmen" in wb.sheetnames
    ws = wb["Maßnahmen"]
    assert "Annahmen!$B$" in str(ws["H4"].value)
    assert str(ws["I4"].value).startswith("=IF(")


def test_excel_funktioniert_auch_ohne_strukturierte_daten():
    """Fallback: nur Antworttext + Markdown-Tabelle → trotzdem verwertbare Datei."""
    antwort = ("## Bilanz\n\n| Verkehrsträger | t CO₂ | Anteil |\n| --- | --- | --- |\n"
               "| Pkw Diesel | 127,6 | 28,9 |\n| Busse | 307,5 | 69,7 |\n")
    daten = exports.build_workbook(tile_name="Argumentationshilfe", calc=None, structured=None, answer=antwort)
    wb = _wb(daten)
    assert "Antworttext" in wb.sheetnames
    assert "Ergebnis" in wb.sheetnames


def test_markdown_tabelle_und_zahlenparser():
    text = "| A | B |\n| --- | --- |\n| 1.234,5 | x |\n"
    zeilen = exports.markdown_tabelle_zu_zeilen(text)
    assert zeilen == [["A", "B"], ["1.234,5", "x"]]
    assert exports._zahlenwert("1.234,5") == 1234.5
    assert exports._zahlenwert("28,9 %") == 28.9
    assert exports._zahlenwert("Diesel-Pkw") == "Diesel-Pkw"


def test_json_artefakt_ist_maschinenlesbar():
    import json

    daten = exports.json_artefakt("CO₂-Bilanz", {"typ": "co2_bilanz", "summe_t_co2": 441.0}, {"annahmen": []})
    obj = json.loads(daten)
    assert obj["rechenkern"]["summe_t_co2"] == 441.0
    assert obj["kachel"] == "CO₂-Bilanz"
