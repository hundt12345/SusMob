"""Tests für Structured Output (Phase 1) und die Exporterzeugung."""
import io

from openpyxl import load_workbook

from server import export, schemas

CO2_DATEN = {
    "titel": "CO₂-Bilanz Fuhrpark 2024",
    "zusammenfassung": "Bilanz auf Basis der Angaben der Kommune.",
    "basisjahr": "2024",
    "bilanz": [
        {"verkehrstraeger": "Pkw Diesel", "fahrleistung_km_jahr": 700000, "energieverbrauch": 47600,
         "energie_einheit": "l", "emissionsfaktor": "2.680 g CO₂/l", "co2_t_jahr": 127.6,
         "anteil_prozent": 29.3, "quelle": "Angabe der Kommune"},
        {"verkehrstraeger": "Bus Diesel", "fahrleistung_km_jahr": 255000, "energieverbrauch": 114750,
         "energie_einheit": "l", "emissionsfaktor": "2.680 g CO₂/l", "co2_t_jahr": 307.5,
         "anteil_prozent": 70.7, "quelle": "Angabe der Kommune"},
    ],
    "summe_t_co2_jahr": 435.1,
    "einsparpotenziale": [
        {"name": "Diesel-Busse → E-Busse", "wirkung": "184 t CO₂/a", "kosten_band": "150–400 k€",
         "prioritaet": "hoch"},
    ],
    "annahmen": ["Diesel 2.680 g CO₂/l", "Strommix 372 g CO₂/kWh"],
    "quellen": ["UBA-Emissionsfaktoren"],
    "naechste_schritte": ["Daten mit dem Bauhof abgleichen"],
}


def test_extract_json_mit_zaun_und_prosa():
    assert schemas.extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert schemas.extract_json('Hier das Ergebnis: {"a": {"b": 2}} fertig.') == {"a": {"b": 2}}
    assert schemas.extract_json("kein json") is None


def test_validate_erkennt_fehler_und_akzeptiert_gueltige_daten():
    obj, err = schemas.validate("co2-bilanz", CO2_DATEN)
    assert err == [] and obj is not None and obj.summe_t_co2_jahr == 435.1
    kaputt = dict(CO2_DATEN)
    kaputt["bilanz"] = "keine Liste"
    obj2, err2 = schemas.validate("co2-bilanz", kaputt)
    assert obj2 is None and err2


def test_schema_ist_json_schema_mit_zusatzfeldern_verboten():
    s = schemas.schema_for("co2-bilanz")
    assert s["additionalProperties"] is False
    assert "summe_t_co2_jahr" in s["properties"]
    # Unbekannte Kachel fällt auf das allgemeine Schema zurück
    assert schemas.model_for("gibt-es-nicht").__name__ == "AllgemeinErgebnis"


def test_to_markdown_enthaelt_tabellen_und_annahmen():
    md = schemas.to_markdown("co2-bilanz", CO2_DATEN)
    assert "| Verkehrsträger |" in md
    assert "**Summe: 435.1 t CO₂/a**" in md
    assert "**Annahmen**" in md
    assert "Prüfen" in md or "prüfen" in md


def test_flat_rows_liefert_blaetter():
    rows = schemas.flat_rows("co2-bilanz", CO2_DATEN)
    assert rows["Bilanz"][0]["verkehrstraeger"] == "Pkw Diesel"
    assert any(k.startswith("Massnahmen") for k in rows)
    assert rows["Naechste_Schritte"][0]["schritt"]


def test_xlsx_hat_annahmen_formeln_und_chart():
    raw = export.build_xlsx("co2-bilanz", {**CO2_DATEN, "eingaben": {"pkw": 25}}, "Test")
    assert len(raw) > 5000
    wb = load_workbook(io.BytesIO(raw))
    assert "Annahmen" in wb.sheetnames
    assert wb["Annahmen"]["B3"].value == 2680  # Diesel-Faktor ist editierbar
    ws = wb["Bilanz"]
    formel = ws.cell(row=2, column=6).value
    assert isinstance(formel, str) and formel.startswith("=") and "Annahmen!$B$" in formel
    assert str(ws.cell(row=4, column=6).value).startswith("=SUM(")
    assert ws._charts, "Balkendiagramm fehlt"
    assert wb["Eingangsdaten"]["A3"].value == "pkw"
    assert wb["Eingangsdaten"]["B3"].value == 25


def test_csv_json_markdown_und_svg():
    csv_text = export.build_csv("co2-bilanz", CO2_DATEN)
    assert csv_text.splitlines()[0] == "verkehrstraeger;fahrleistung_km_jahr;energieverbrauch;energie_einheit;emissionsfaktor;co2_t_jahr;anteil_prozent;quelle"
    js = export.build_json("co2-bilanz", CO2_DATEN)
    assert '"summe_t_co2_jahr": 435.1' in js
    diagramme = dict(export.diagram_for("co2-bilanz", CO2_DATEN))
    assert diagramme["bilanz_donut.svg"].startswith("<svg")
    assert "Bus Diesel" in diagramme["bilanz_donut.svg"]


def test_print_html_ist_druckfaehig():
    html = export.build_print_html("co2-bilanz", CO2_DATEN, "CO₂-Bilanz 2024", "Musterstadt")
    assert "@page" in html and "size: A4" in html
    assert "SusMob" in html and "Musterstadt" in html
    assert "<table>" in html
    assert "Als PDF speichern" in html


def test_store_and_load_artifact(tmp_path):
    a = export.store_artifact(conversation_id="c1", tile_id="co2-bilanz", kind="json",
                              filename="t.json", content='{"a":1}', payload={"x": 1},
                              artifacts_dir=tmp_path)
    assert a["size"] == 7
    path, name = export.artifact_bytes(a["id"])
    assert name == "t.json" and path.read_text() == '{"a":1}'
