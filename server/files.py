"""Text-Extraktion aus hochgeladenen Dateien (für den LLM-Kontext) + Bild-Erkennung."""
from __future__ import annotations

import base64
from pathlib import Path

MAX_FILE_CHARS = 12_000  # pro Datei
TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".log", ".json"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
MIME_BY_EXT = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp",
}


def is_image(path: Path, mime: str = "") -> bool:
    """Bilddateien werden nicht text-extraniert, sondern multimodal ans Modell gegeben (Plan 4.4)."""
    ext = Path(path).suffix.lower()
    return ext in IMAGE_EXTS or mime.startswith("image/")


def image_data_url(path: Path, mime: str = "") -> str:
    """Bild als Data-URL (OpenRouter/OpenAI-Format für Vision-Modelle)."""
    p = Path(path)
    typ = mime or MIME_BY_EXT.get(p.suffix.lower(), "image/png")
    return f"data:{typ};base64,{base64.b64encode(p.read_bytes()).decode('ascii')}"


def _truncate(text: str) -> str:
    if len(text) <= MAX_FILE_CHARS:
        return text
    return text[:MAX_FILE_CHARS] + f"\n… (Datei gekürzt auf {MAX_FILE_CHARS} Zeichen)"


def extract_text(path: Path) -> str:
    """Gibt lesbaren Text aus der Datei zurück; leer, wenn nicht extrahierbar."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext in TEXT_EXTS:
            text = p.read_text(encoding="utf-8", errors="replace")
            return _truncate(text.strip())

        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(p, read_only=True, data_only=True)
            out: list[str] = []
            for ws in wb.worksheets:
                out.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(v.strip() for v in vals):
                        out.append(" | ".join(vals))
                if len("\n".join(out)) > MAX_FILE_CHARS * 2:
                    break
            wb.close()
            return _truncate("\n".join(out))

        if ext == ".docx":
            from docx import Document

            doc = Document(p)
            out = [p_.text for p_ in doc.paragraphs if p_.text.strip()]
            for tbl in doc.tables:
                for r in tbl.rows:
                    out.append(" | ".join(c.text.strip() for c in r.cells))
            return _truncate("\n".join(out))

        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(p)
            return _truncate("\n".join((pg.extract_text() or "") for pg in reader.pages))
    except Exception:
        return ""
    return ""


SUPPORTED = " .txt .md .csv .tsv .log .json .xlsx .docx .pdf .png .jpg .jpeg .webp .gif .bmp "
